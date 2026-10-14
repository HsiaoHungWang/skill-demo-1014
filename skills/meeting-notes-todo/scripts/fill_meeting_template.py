#!/usr/bin/env python3
"""
fill_meeting_template.py — 把整理好的會議摘要與待辦事項，填進 iSpan 官方會議紀錄 Word範本。

用法：
    python fill_meeting_template.py data.json output.docx

data.json 格式：
{
  "meeting_type": "月會",              // 營管會議 / 月會 / 外部會議 / 其他
  "date": "2026-09-14",
  "location": "教育訓練部會議室",
  "meeting_name": "教育訓練部門週會",
  "chair": "王雅婷",
  "attendees": "王雅婷、徐建宏、陳雅廷",
  "recorder": "徐建宏",
  "decisions": [                        // 對應「會議紀錄及決議事項」，範本預設9列，超過會自動加列
    "Q4 AI應用課程開班時間確認為11月中（原定10月底，因兩位候選講師檔期無法配合，修正為11月中）。",
    "課程大綱初稿已完成，將請R&D同仁協助確認內容正確性後定稿。"
  ],
  "actions": [                          // 對應「會議待執行事項」，範本預設3列，超過會自動加列
    {"item": "傳送AI應用課程大綱初稿給大家", "owner": "小徐", "due": "2026-09-16"},
    {"item": "完成招生頁面文案", "owner": "阿廷", "due": "2026-09-18"}
  ]
}

範本本身（assets/meeting-minutes-template.docx）完全不需要修改或加入佔位符——
這支script是用表格的「列/欄位置」直接定位儲存格並填值，
這樣可以完整保留原本的公司抬頭、框線、自動編號等格式，不會因為打包/解壓縮而遺失。
項目數超過範本預設列數時，會複製最後一列（含框線、字型、自動編號）往下加列，範本檔本身不會被修改。
"""

import sys
import json
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn

TEMPLATE_PATH_DEFAULT = "assets/meeting-minutes-template.docx"
MEETING_TYPES = ["營管會議", "月會", "外部會議", "其他"]
DEFAULT_DECISION_ROWS = 9  # 範本預設列數（不是上限）
DEFAULT_ACTION_ROWS = 3


def ensure_rows(table, needed):
    """確保表格至少有 needed 列；不夠就複製最後一列加在後面，保留原本格式。"""
    added = 0
    while len(table.rows) < needed:
        last_tr = table.rows[-1]._tr
        new_tr = deepcopy(last_tr)
        for t in new_tr.iter(qn("w:t")):
            t.text = ""
        last_tr.addnext(new_tr)
        added += 1
    return added


def set_cell_text(cell, text):
    """在儲存格既有的第一個段落加入文字，保留原本的段落格式（含自動編號）。"""
    p = cell.paragraphs[0]
    # 清掉可能存在的空白 run，避免重複填值時疊加
    for run in list(p.runs):
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def check_meeting_type(doc, meeting_type):
    """把對應會議類型前面的 □ 換成 ■。

    範本裡每個 □ 都是獨立的一個 run，後面接著 1 個以上的 run 拼出該選項的
    文字（例如「營管」+「會議」），直到下一個 □ 或段落結尾為止。
    """
    if meeting_type not in MEETING_TYPES:
        return

    target_para = doc.paragraphs[0]
    runs = target_para.runs

    box_indices = [i for i, r in enumerate(runs) if r.text == "□"]
    for idx, box_i in enumerate(box_indices):
        next_box_i = box_indices[idx + 1] if idx + 1 < len(box_indices) else len(runs)
        label = "".join(r.text for r in runs[box_i + 1:next_box_i]).strip()
        if label.startswith(meeting_type):
            runs[box_i].text = "■"
            return


def fill_header_table(doc, data):
    """table[0]：日期／地點／會議名稱／會議主席／與會人員／記錄"""
    t = doc.tables[0]
    set_cell_text(t.rows[0].cells[1], data.get("date", ""))
    set_cell_text(t.rows[0].cells[3], data.get("location", ""))
    set_cell_text(t.rows[1].cells[1], data.get("meeting_name", ""))
    set_cell_text(t.rows[2].cells[1], data.get("chair", ""))
    set_cell_text(t.rows[3].cells[1], data.get("attendees", ""))
    set_cell_text(t.rows[4].cells[1], data.get("recorder", ""))


def fill_decisions_table(doc, decisions):
    """table[1]：會議紀錄及決議事項，範本預設9列，自動編號已存在於範本中；超過9項自動加列。"""
    t = doc.tables[1]
    added = ensure_rows(t, len(decisions))
    for i, text in enumerate(decisions):
        set_cell_text(t.rows[i].cells[0], text)
    if added:
        print(f"決議事項有 {len(decisions)} 項，已自動新增 {added} 列。", file=sys.stderr)


def fill_actions_table(doc, actions):
    """table[2]：會議待執行事項，row 0 是欄位標題，範本預設3列資料；超過3項自動加列。
    欄位為 項目/負責人員/完成日期。"""
    t = doc.tables[2]
    added = ensure_rows(t, len(actions) + 1)  # +1 是標題列
    for i, action in enumerate(actions):
        row = t.rows[i + 1]
        set_cell_text(row.cells[0], action.get("item", ""))
        set_cell_text(row.cells[1], action.get("owner", ""))
        set_cell_text(row.cells[2], action.get("due") or "待定")
    if added:
        print(f"待辦事項有 {len(actions)} 項，已自動新增 {added} 列。", file=sys.stderr)


def main():
    if len(sys.argv) not in (3, 4):
        print(
            "用法：python fill_meeting_template.py data.json output.docx [template.docx]",
            file=sys.stderr,
        )
        sys.exit(1)

    data_path = sys.argv[1]
    output_path = sys.argv[2]
    template_path = sys.argv[3] if len(sys.argv) == 4 else TEMPLATE_PATH_DEFAULT

    with open(data_path, encoding="utf-8") as f:
        data = json.load(f)

    doc = Document(template_path)

    check_meeting_type(doc, data.get("meeting_type", ""))
    fill_header_table(doc, data)
    fill_decisions_table(doc, data.get("decisions", []))
    fill_actions_table(doc, data.get("actions", []))

    doc.save(output_path)
    print(f"已產生 {output_path}")


if __name__ == "__main__":
    main()
